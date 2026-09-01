# Regressionstests fuer v0.9.8:
#   * eingebettete BIP-39-Wortliste ("Language not detected" im Onefile-Build)
#   * ETH-Ableitung ohne eth_account-Wortlisten (bitidentisch zu vorher)
#   * Saldo-Anzeige bei Netzwerkfehlern (Cache-Fallback statt 0)
#   * TronGrid-Pagination
#   * Transaktions-Export (CSV/Excel)
#
# Alle Tests laufen offline (kein Netzwerk, keine Wallet-Datei noetig).
# Ausfuehren:  python -m unittest tests.test_v098 -v

import csv
import os
import sys
import tempfile
import unittest
import warnings
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
warnings.simplefilter("ignore", DeprecationWarning)

from src.wallet.bip39_wordlist import ENGLISH, get_mnemonic, verify_wordlist
from src.wallet.doi_wallet import DoiWallet
from src.wallet.seed_manager import SeedManager
from src.wallet.tron_network import TronClient
from src.utils import tx_export

TEST_MNEMONIC = (
    "abandon abandon abandon abandon abandon abandon abandon abandon "
    "abandon abandon abandon about"
)
# Bekannter Testvektor (Trezor / eth_account) fuer m/44'/60'/0'/0/0
ETH_REF_ADDRESS = "0x9858EfFD232B4033E47d90003D41EC34EcaEda94"


class TestEmbeddedWordlist(unittest.TestCase):
    """Befund: Mnemonic("english") las die Wortliste bei jedem Aufruf von
    der Platte; im Onefile-Build verschwand sie aus %TEMP% -> "Language not
    detected" beim Oeffnen weiterer Wallets."""

    def test_wordlist_integrity(self):
        self.assertEqual(len(ENGLISH), 2048)
        self.assertTrue(verify_wordlist())
        self.assertEqual(ENGLISH[0], "abandon")
        self.assertEqual(ENGLISH[-1], "zoo")

    def test_get_mnemonic_is_cached_and_works_without_files(self):
        m1 = get_mnemonic()
        m2 = get_mnemonic()
        self.assertIs(m1, m2)
        self.assertTrue(m1.check(TEST_MNEMONIC))
        self.assertFalse(m1.check("abandon abandon zzz"))

    def test_mnemonic_works_even_if_package_wordlist_dir_missing(self):
        # Simuliert das geloeschte Temp-Verzeichnis: os.path.exists liefert
        # fuer die Paket-Wortliste False. get_mnemonic() darf davon nichts merken.
        import mnemonic as mn_pkg
        real_exists = os.path.exists

        def fake_exists(p):
            if "wordlist" in str(p):
                return False
            return real_exists(p)

        with mock.patch("mnemonic.mnemonic.os.path.exists", side_effect=fake_exists):
            with self.assertRaises(Exception):
                mn_pkg.Mnemonic("english")          # Altes Verhalten: Fehler
            sm = SeedManager().from_mnemonic(TEST_MNEMONIC)   # Neues Verhalten: ok
            self.assertTrue(sm.get_receive_address(0).startswith("N"))


class TestEthDerivation(unittest.TestCase):
    """Befund: Account.from_mnemonic() braucht Wortlisten-Dateien und
    schreibt bei Fehlern die Seed-Woerter in die Fehlermeldung."""

    def setUp(self):
        try:
            from src.wallet.eth_wallet import EthWallet
        except ImportError:
            self.skipTest("web3/eth_account nicht installiert")
        self.EthWallet = EthWallet

    def test_reference_vector(self):
        self.assertEqual(self.EthWallet().from_mnemonic(TEST_MNEMONIC), ETH_REF_ADDRESS)

    def test_identical_to_eth_account_for_random_seeds(self):
        from eth_account import Account
        Account.enable_unaudited_hdwallet_features()
        for idx in (0, 1, 7):
            m = get_mnemonic().generate(256)
            ref = Account.from_mnemonic(m, account_path=f"m/44'/60'/0'/0/{idx}").address
            self.assertEqual(self.EthWallet().from_mnemonic(m, idx), ref)

    def test_error_message_does_not_leak_words(self):
        bad = "abandon abandon zzz"
        with self.assertRaises(ValueError) as ctx:
            self.EthWallet().from_mnemonic(bad)
        self.assertNotIn("zzz", str(ctx.exception))
        self.assertNotIn("abandon", str(ctx.exception))


class _FakeElectrum:
    """ElectrumX-Attrappe: bestimmte Adressen liefern Fehler."""

    def __init__(self, balances, fail=(), drop_connection_on=None):
        self.balances = balances
        self.fail = set(fail)
        self.drop_on = drop_connection_on
        self.is_connected = True
        self.calls = []

    def get_balance(self, addr):
        self.calls.append(addr)
        if addr == self.drop_on:
            self.is_connected = False
            raise ConnectionError("Verbindung verloren")
        if addr in self.fail:
            raise ConnectionError("timeout")
        return dict(self.balances.get(addr, {"confirmed": 0, "unconfirmed": 0}))


class TestBalanceFallback(unittest.TestCase):
    """Befund: Nach einem Netzwerkfehler fehlten Adressen komplett in der
    Summe -> Dashboard zeigte 0 DOI, Info-Dialog (Cache) den echten Saldo."""

    def _wallet(self):
        w = DoiWallet()
        w.restore(TEST_MNEMONIC)
        return w

    def test_cached_value_is_used_when_query_fails(self):
        w = self._wallet()
        addrs = w.get_all_addresses()[:3]
        a0, a1, a2 = addrs
        # Erste Abfrage: alles ok, Cache fuellen
        w.electrum = _FakeElectrum({a0: {"confirmed": 500, "unconfirmed": 0},
                                    a1: {"confirmed": 700, "unconfirmed": 0}})
        # Nur die drei Adressen betrachten (Test bleibt schnell)
        w._known_addresses = {a: w._known_addresses[a] for a in addrs}
        first = w.get_balance(force_refresh=True)
        self.assertEqual(first["confirmed"], 1200)
        self.assertTrue(first["complete"])

        # Zweite Abfrage: a1 faellt aus -> Cache-Wert bleibt in der Summe
        w.electrum = _FakeElectrum({a0: {"confirmed": 500, "unconfirmed": 0}}, fail=(a1,))
        second = w.get_balance(force_refresh=True)
        self.assertEqual(second["confirmed"], 1200)
        self.assertFalse(second["complete"])
        self.assertEqual(second["stale_addresses"], [a1])
        self.assertEqual(second["unknown_addresses"], [])

    def test_unknown_when_no_cache(self):
        w = self._wallet()
        addrs = w.get_all_addresses()[:2]
        w._known_addresses = {a: w._known_addresses[a] for a in addrs}
        w.electrum = _FakeElectrum({}, fail=(addrs[1],))
        res = w.get_balance(force_refresh=True)
        self.assertFalse(res["complete"])
        self.assertEqual(res["unknown_addresses"], [addrs[1]])

    def test_abort_after_connection_loss(self):
        w = self._wallet()
        addrs = w.get_all_addresses()[:5]
        w._known_addresses = {a: w._known_addresses[a] for a in addrs}
        w.electrum = _FakeElectrum({}, drop_connection_on=addrs[1])
        res = w.get_balance(force_refresh=True)
        # Nach dem Verbindungsverlust keine weiteren Abfragen mehr
        self.assertEqual(w.electrum.calls, addrs[:2])
        self.assertEqual(len(res["stale_addresses"]), 4)
        self.assertFalse(res["complete"])


class TestTronPagination(unittest.TestCase):
    """Befund: Nur die erste TronGrid-Seite (20 Eintraege) wurde geladen."""

    def _client(self, pages):
        client = TronClient.__new__(TronClient)
        client.network = {"usdt_contract": "TXYZ"}
        calls = []

        def fake_get(endpoint, params=None, timeout=15):
            calls.append(dict(params or {}))
            idx = len(calls) - 1
            data, fp = pages[idx] if idx < len(pages) else ([], None)
            out = {"data": data, "meta": {}}
            if fp:
                out["meta"]["fingerprint"] = fp
            return out

        client._get = fake_get
        client._calls = calls
        return client

    def test_follows_fingerprint_until_end(self):
        pages = [([{"txID": f"a{i}"} for i in range(200)], "fp1"),
                 ([{"txID": f"b{i}"} for i in range(200)], "fp2"),
                 ([{"txID": "c0"}], None)]
        c = self._client(pages)
        result = c.get_transactions("TADDR", limit=None)
        self.assertEqual(len(result), 401)
        self.assertEqual(len(c._calls), 3)
        self.assertNotIn("fingerprint", c._calls[0])
        self.assertEqual(c._calls[1]["fingerprint"], "fp1")
        self.assertEqual(c._calls[2]["fingerprint"], "fp2")

    def test_limit_respected(self):
        pages = [([{"txID": f"a{i}"} for i in range(200)], "fp1"),
                 ([{"txID": f"b{i}"} for i in range(200)], "fp2")]
        c = self._client(pages)
        result = c.get_transactions("TADDR", limit=250)
        self.assertEqual(len(result), 250)
        self.assertEqual(len(c._calls), 2)
        self.assertEqual(c._calls[1]["limit"], 50)

    def test_default_single_page(self):
        pages = [([{"txID": f"a{i}"} for i in range(20)], "fp1")]
        c = self._client(pages)
        self.assertEqual(len(c.get_trc20_transactions("TADDR")), 20)
        self.assertEqual(len(c._calls), 1)
        self.assertEqual(c._calls[0]["contract_address"], "TXYZ")


class TestTxExport(unittest.TestCase):
    TXS = [
        {"hash": "aa", "direction": "received", "value": 12.5, "symbol": "DOI",
         "timestamp": 1_720_000_000, "from": "", "to": "NAbc", "block": 100},
        {"hash": "bb", "direction": "sent", "value": 0.25, "symbol": "DOI",
         "timestamp": 1_720_100_000, "from": "", "to": "", "block": 0},
        {"hash": "cc", "direction": "received", "value": 3, "symbol": "TRX",
         "timestamp": 1_720_200_000, "from": "TA", "to": "TB", "block": 5},
        {"hash": "aa", "direction": "received", "value": 12.5, "symbol": "DOI",
         "timestamp": 1_720_000_000, "from": "", "to": "NAbc", "block": 100},
    ]

    def test_build_rows_filters_symbol_dedups_and_sorts(self):
        rows = tx_export.build_rows("rth dat", "DOI", self.TXS, notes={"aa": "Kauf"})
        self.assertEqual([r["TX-Hash"] for r in rows], ["bb", "aa"])
        self.assertEqual(rows[1]["Notiz"], "Kauf")
        self.assertEqual(rows[0]["Richtung"], "Ausgang")
        self.assertEqual(rows[0]["Betrag_signiert"], -0.25)
        self.assertEqual(rows[0]["Status"], "Unbestaetigt")
        self.assertEqual(rows[1]["Status"], "Bestaetigt")
        self.assertEqual(rows[1]["Wallet"], "rth dat")
        self.assertEqual(rows[1]["Waehrung"], "DOI")

    def test_csv_is_excel_compatible(self):
        rows = tx_export.build_rows("rth dat", "DOI", self.TXS)
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / tx_export.export_filename("rth dat", "DOI", "csv")
            tx_export.write_csv(path, rows)
            raw = path.read_bytes()
            self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))      # UTF-8-BOM
            self.assertIn(b"\r\n", raw)                          # CRLF
            with open(path, encoding="utf-8-sig", newline="") as f:
                data = list(csv.reader(f, delimiter=";"))
            self.assertEqual(data[0], tx_export.COLUMNS)
            self.assertEqual(len(data), 3)
            betrag_idx = tx_export.COLUMNS.index("Betrag_signiert")
            self.assertEqual(data[1][betrag_idx], "-0,25")       # Dezimalkomma
            self.assertEqual(data[2][betrag_idx], "12,5")

    def test_filename_and_sheetname_sanitized(self):
        self.assertNotIn("/", tx_export.safe_filename("a/b:c*d"))
        used = set()
        n1 = tx_export.safe_sheet_name("Wallet mit [sehr] langem Namen: DOI/USDT", used)
        n2 = tx_export.safe_sheet_name("Wallet mit [sehr] langem Namen: DOI/USDT", used)
        self.assertLessEqual(len(n1), 31)
        self.assertLessEqual(len(n2), 31)
        self.assertNotEqual(n1, n2)
        for bad in "[]:*?/\\":
            self.assertNotIn(bad, n1)

    @unittest.skipUnless(tx_export.HAS_OPENPYXL, "openpyxl nicht installiert")
    def test_xlsx_roundtrip(self):
        import openpyxl
        sheets = {
            "rth dat – DOI": tx_export.build_rows("rth dat", "DOI", self.TXS),
            "rth dat – TRX": tx_export.build_rows("rth dat", "TRX", self.TXS),
        }
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "export.xlsx"
            tx_export.write_xlsx(path, sheets)
            wb = openpyxl.load_workbook(path)
            self.assertEqual(len(wb.sheetnames), 2)
            ws = wb[wb.sheetnames[0]]
            self.assertEqual([c.value for c in ws[1]], tx_export.COLUMNS)
            self.assertEqual(ws.max_row, 3)
            self.assertAlmostEqual(ws.cell(row=2, column=7).value, -0.25)


if __name__ == "__main__":
    unittest.main()

"""
Transaktions-Export (CSV / Excel) – v0.9.8

Reine Datei-/Formatlogik ohne GUI- oder Netzwerkabhaengigkeit, damit sie
offline testbar bleibt. Die GUI (wallet_gui.py, ExportDialog) sammelt die
normalisierten Transaktionen je Wallet und Waehrung und ruft hier
build_rows() + write_csv() / write_xlsx() auf.

Eingabeformat je Transaktion (normalisiert, wie in WalletApp._normalize_tx):
    {
        "hash": str, "direction": "received"|"sent"|"unknown",
        "value": float, "symbol": "DOI"|"TRX"|"USDT"|"ETH"|"wDOI",
        "timestamp": int (Unix, 0 = unbekannt), "from": str, "to": str,
        "block": int (0 = unbestaetigt), "note": str (optional, z.B. "gas")
    }

CSV-Format: Semikolon-getrennt, UTF-8 mit BOM, Dezimalkomma, CRLF –
oeffnet sich in deutschem Excel per Doppelklick korrekt.
Excel-Format: openpyxl (optional; ohne openpyxl nur CSV).

© 2026 Ottmar Neuburger, WEBanizer AG – MIT License
"""

import csv
import re
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:  # pragma: no cover - haengt von der Umgebung ab
    openpyxl = None
    HAS_OPENPYXL = False

SYMBOLS = ("DOI", "TRX", "USDT", "ETH", "wDOI")

# Nachkommastellen je Waehrung (volle Genauigkeit fuer die Buchhaltung)
DECIMALS = {"DOI": 8, "TRX": 6, "USDT": 6, "ETH": 8, "wDOI": 8}

COLUMNS = [
    "Datum", "Zeit", "Wallet", "Waehrung", "Richtung", "Betrag",
    "Betrag_signiert", "Von", "An", "Block", "Status", "TX-Hash", "Notiz",
]

DIRECTION_LABEL = {
    "received": "Eingang",
    "sent": "Ausgang",
    "unknown": "Unbekannt",
}


def format_amount(value, symbol: str) -> str:
    """Betrag als Zeichenkette mit Dezimalpunkt und fester Stellenzahl."""
    try:
        v = float(value or 0)
    except (TypeError, ValueError):
        v = 0.0
    return f"{v:.{DECIMALS.get(symbol, 8)}f}"


def _ts_to_date_time(ts) -> tuple:
    try:
        ts = int(ts or 0)
    except (TypeError, ValueError):
        ts = 0
    if ts <= 0:
        return "", ""
    try:
        dt = datetime.fromtimestamp(ts)
    except (OverflowError, OSError, ValueError):
        return "", ""
    return dt.strftime("%d.%m.%Y"), dt.strftime("%H:%M:%S")


def tx_to_row(tx: dict, wallet_name: str, notes: Optional[dict] = None) -> dict:
    """Wandelt eine normalisierte Transaktion in eine Export-Zeile um."""
    symbol = str(tx.get("symbol", "?"))
    direction = str(tx.get("direction", "unknown"))
    value = tx.get("value", 0) or 0
    try:
        value = float(value)
    except (TypeError, ValueError):
        value = 0.0
    signed = -abs(value) if direction == "sent" else abs(value)
    if direction == "unknown":
        signed = 0.0
    date_str, time_str = _ts_to_date_time(tx.get("timestamp", 0))
    try:
        block = int(tx.get("block", tx.get("height", 0)) or 0)
    except (TypeError, ValueError):
        block = 0
    tx_hash = str(tx.get("hash", ""))
    note = ""
    if notes and tx_hash in notes:
        note = str(notes[tx_hash])
    elif tx.get("note"):
        note = str(tx.get("note"))
        if note == "gas":
            note = "Gas-Gebuehr"

    return {
        "Datum": date_str,
        "Zeit": time_str,
        "Wallet": wallet_name,
        "Waehrung": symbol,
        "Richtung": DIRECTION_LABEL.get(direction, direction),
        "Betrag": abs(value),
        "Betrag_signiert": signed,
        "Von": str(tx.get("from", "") or ""),
        "An": str(tx.get("to", "") or ""),
        "Block": block,
        "Status": "Bestaetigt" if block > 0 else "Unbestaetigt",
        "TX-Hash": tx_hash,
        "Notiz": note,
    }


def build_rows(wallet_name: str, symbol: str, txs: Iterable[dict],
               notes: Optional[dict] = None) -> List[dict]:
    """
    Erzeugt die Export-Zeilen fuer ein Wallet und eine Waehrung.

    Transaktionen anderer Waehrungen werden ignoriert, Dubletten (gleicher
    Hash + gleiche Notiz) entfernt, Sortierung neueste zuerst.
    """
    seen = set()
    rows = []
    for tx in txs or []:
        if not isinstance(tx, dict):
            continue
        if str(tx.get("symbol", "")) != symbol:
            continue
        key = (tx.get("hash", ""), tx.get("note", ""))
        if key in seen:
            continue
        seen.add(key)
        rows.append(tx_to_row(tx, wallet_name, notes))
    rows.sort(key=lambda r: (r["Datum"] and datetime.strptime(
        r["Datum"] + " " + r["Zeit"], "%d.%m.%Y %H:%M:%S").timestamp() or 0),
        reverse=True)
    return rows


def safe_filename(text: str, max_len: int = 60) -> str:
    """Entfernt Zeichen, die in Dateinamen nicht erlaubt sind."""
    cleaned = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", str(text)).strip(" ._")
    cleaned = re.sub(r"\s+", " ", cleaned)
    return (cleaned or "Wallet")[:max_len]


def safe_sheet_name(text: str, used: Optional[set] = None) -> str:
    """Excel-Blattname: max. 31 Zeichen, ohne []:*?/\\ und eindeutig."""
    cleaned = re.sub(r"[\[\]:*?/\\]+", "_", str(text)).strip() or "Blatt"
    base = cleaned[:31]
    if used is None:
        return base
    name = base
    n = 2
    while name in used:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def _csv_cell(value, column: str) -> str:
    if column in ("Betrag", "Betrag_signiert"):
        # Dezimalkomma fuer deutsches Excel
        return f"{float(value):.8f}".rstrip("0").rstrip(".").replace(".", ",") or "0"
    return "" if value is None else str(value)


def write_csv(path, rows: List[dict]) -> str:
    """Schreibt Zeilen als Excel-kompatible CSV (UTF-8-BOM, Semikolon)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";", quoting=csv.QUOTE_MINIMAL,
                            lineterminator="\r\n")
        writer.writerow(COLUMNS)
        for row in rows:
            writer.writerow([_csv_cell(row.get(c, ""), c) for c in COLUMNS])
    return str(path)


def write_xlsx(path, sheets: Dict[str, List[dict]]) -> str:
    """
    Schreibt eine Excel-Arbeitsmappe mit einem Blatt je Eintrag in `sheets`.

    Args:
        path: Zieldatei (.xlsx)
        sheets: {Blattname: Zeilen} – Reihenfolge bleibt erhalten.

    Raises:
        RuntimeError: wenn openpyxl nicht installiert ist.
    """
    if not HAS_OPENPYXL:
        raise RuntimeError("Excel-Export benoetigt das Paket openpyxl "
                           "(pip install openpyxl). CSV-Export ist verfuegbar.")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used_names: set = set()
    if not sheets:
        sheets = {"Transaktionen": []}

    for title, rows in sheets.items():
        ws = wb.create_sheet(safe_sheet_name(title, used_names))
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            ws.append([row.get(c, "") for c in COLUMNS])
        # Zahlenformat + Spaltenbreiten
        amount_cols = [COLUMNS.index("Betrag") + 1, COLUMNS.index("Betrag_signiert") + 1]
        for r in range(2, ws.max_row + 1):
            for c in amount_cols:
                ws.cell(row=r, column=c).number_format = "0.00000000"
        widths = {
            "Datum": 12, "Zeit": 10, "Wallet": 18, "Waehrung": 10,
            "Richtung": 11, "Betrag": 18, "Betrag_signiert": 18, "Von": 44,
            "An": 44, "Block": 10, "Status": 13, "TX-Hash": 66, "Notiz": 30,
        }
        for i, col in enumerate(COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 14)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb.save(path)
    return str(path)


def export_filename(wallet_name: str, symbol: str, ext: str,
                    when: Optional[datetime] = None) -> str:
    """Dateiname `Transaktionen_<Wallet>_<Waehrung>_<JJJJ-MM-TT>.<ext>`."""
    when = when or datetime.now()
    stamp = when.strftime("%Y-%m-%d")
    return f"Transaktionen_{safe_filename(wallet_name)}_{symbol}_{stamp}.{ext}"
